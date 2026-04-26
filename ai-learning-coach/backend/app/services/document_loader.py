"""
Multi-format Document Loader for RAG pipeline.

Supports: .pdf, .docx, .txt, .md files.
Handles chunking and metadata extraction.
"""

import os
from typing import List, Tuple, Dict, Any
from flask import current_app

# Try to import optional document processing libraries
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None


def _chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
    """Split text into overlapping chunks."""
    if not text:
        return []
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        chunks.append(text[start:end])
        start = end - overlap if end < length else end
    return chunks


def load_pdf(file_path: str, chunk_size: int = 1000, overlap: int = 200) -> Tuple[List[str], Dict[str, Any]]:
    """
    Load and chunk a PDF file.
    
    Args:
        file_path: Path to PDF file
        chunk_size: Characters per chunk
        overlap: Overlap between chunks
        
    Returns:
        Tuple of (chunks, metadata)
    """
    if PdfReader is None:
        raise RuntimeError("PyPDF2 not installed; install with: pip install PyPDF2")
    
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"PDF file not found: {file_path}")
    
    try:
        with open(file_path, "rb") as fh:
            reader = PdfReader(fh)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
        
        chunks = _chunk_text(text, chunk_size=chunk_size, overlap=overlap)
        metadata = {
            "source": os.path.basename(file_path),
            "file_type": "pdf",
            "total_pages": len(reader.pages),
            "total_chars": len(text),
        }
        return chunks, metadata
    except Exception as exc:
        current_app.logger.error("Failed to load PDF %s: %s", file_path, exc)
        raise


def load_docx(file_path: str, chunk_size: int = 1000, overlap: int = 200) -> Tuple[List[str], Dict[str, Any]]:
    """
    Load and chunk a DOCX file.
    
    Args:
        file_path: Path to DOCX file
        chunk_size: Characters per chunk
        overlap: Overlap between chunks
        
    Returns:
        Tuple of (chunks, metadata)
    """
    if DocxDocument is None:
        raise RuntimeError("python-docx not installed; install with: pip install python-docx")
    
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")
    
    try:
        doc = DocxDocument(file_path)
        text = "\n".join(para.text for para in doc.paragraphs if para.text)
        
        chunks = _chunk_text(text, chunk_size=chunk_size, overlap=overlap)
        metadata = {
            "source": os.path.basename(file_path),
            "file_type": "docx",
            "total_paragraphs": len(doc.paragraphs),
            "total_chars": len(text),
        }
        return chunks, metadata
    except Exception as exc:
        current_app.logger.error("Failed to load DOCX %s: %s", file_path, exc)
        raise


def load_text(file_path: str, chunk_size: int = 1000, overlap: int = 200) -> Tuple[List[str], Dict[str, Any]]:
    """
    Load and chunk a plain text file (.txt, .md).
    
    Args:
        file_path: Path to text file
        chunk_size: Characters per chunk
        overlap: Overlap between chunks
        
    Returns:
        Tuple of (chunks, metadata)
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Text file not found: {file_path}")
    
    try:
        with open(file_path, "r", encoding="utf-8") as fh:
            text = fh.read()
        
        chunks = _chunk_text(text, chunk_size=chunk_size, overlap=overlap)
        ext = os.path.splitext(file_path)[1].lower()
        metadata = {
            "source": os.path.basename(file_path),
            "file_type": "text" if ext == ".txt" else "markdown",
            "total_chars": len(text),
            "total_lines": len(text.split("\n")),
        }
        return chunks, metadata
    except Exception as exc:
        current_app.logger.error("Failed to load text file %s: %s", file_path, exc)
        raise


def load_document(file_path: str, chunk_size: int = 1000, overlap: int = 200) -> Tuple[List[str], Dict[str, Any]]:
    """
    Auto-detect and load document based on file extension.
    
    Args:
        file_path: Path to document file
        chunk_size: Characters per chunk
        overlap: Overlap between chunks
        
    Returns:
        Tuple of (chunks, metadata)
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".pdf":
        return load_pdf(file_path, chunk_size=chunk_size, overlap=overlap)
    elif ext == ".docx":
        return load_docx(file_path, chunk_size=chunk_size, overlap=overlap)
    elif ext in (".txt", ".md"):
        return load_text(file_path, chunk_size=chunk_size, overlap=overlap)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Supported: .pdf, .docx, .txt, .md")


def load_from_directory(dir_path: str, chunk_size: int = 1000, overlap: int = 200, 
                       recursive: bool = True) -> Tuple[List[Tuple[str, str, Dict]], int]:
    """
    Load all supported documents from a directory.
    
    Args:
        dir_path: Directory path
        chunk_size: Characters per chunk
        overlap: Overlap between chunks
        recursive: Whether to search subdirectories
        
    Returns:
        Tuple of (list of (chunk, source_file, metadata), total_chunks)
    """
    if not os.path.isdir(dir_path):
        raise ValueError(f"Directory not found: {dir_path}")
    
    results = []
    total = 0
    
    pattern = "**/*" if recursive else "*"
    supported_ext = {".pdf", ".docx", ".txt", ".md"}
    
    for root, dirs, files in os.walk(dir_path):
        for fn in files:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in supported_ext:
                continue
            
            full_path = os.path.join(root, fn)
            try:
                chunks, metadata = load_document(full_path, chunk_size=chunk_size, overlap=overlap)
                for chunk in chunks:
                    results.append((chunk, fn, metadata))
                    total += 1
            except Exception as exc:
                current_app.logger.warning("Failed to load document %s: %s", full_path, exc)
                continue
        
        if not recursive:
            break
    
    return results, total
